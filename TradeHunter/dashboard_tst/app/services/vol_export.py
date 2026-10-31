"""The Volatility Surface and the Interpolated Volatility Surface as an Excel workbook (v4.232).

The user (2026-10-31): *"I need to be able to download data from vol surface IV, price, contract expiry, strike (no
need greeks) in to excel. Both the vol surface and the Interpolated"*.

Built from the SAME grid the page draws (``vol_surface.surface`` / ``vol_cm.surface_cm`` with the page's settings),
so the figures are the page's. The page's "Hide new and thinly listed contracts" tidy-up is a view on the page only:
the workbook carries every strike and expiry of the grid.

Sheets
------
* the listed surface - **Contracts** (one row per cell: expiry, days, monthly, strike, call / put, IV %, price =
  Black-Scholes from that IV, last trade) and **IV grid** (strikes down, expiries across, as on screen);
* the interpolated surface - **Contracts** (one row per strike x horizon: horizon days and date, strike, call / put,
  IV %, price at that IV, how the IV was got: a listed expiry, or interpolated between two) and **IV grid
  (moneyness)** / **IV grid (strike)** (the two row modes of the tab).
Each sheet opens with the ticker, the day, the stock price and its source, and the method. No greeks (the user's
choice). openpyxl writes the file (``app/requirements.txt``).
"""
from __future__ import annotations

import datetime as _dt
import io
import re

SIDE_WORDS = {"otm": "out of the money (puts below the price, calls above)", "call": "calls", "put": "puts",
              "mid": "call + put average"}
SPOT_WORDS = {"close": "official close", "parity": "estimated from options", "massive": "stock quote, 15-min delayed"}


def _right(side: str, k: float, spot) -> str:
    if side == "call":
        return "Call"
    if side == "put":
        return "Put"
    if side == "mid":
        return "Call + put average"
    return "Put" if (spot is not None and k < float(spot)) else "Call"


def _num(v, nd=None):
    if v is None:
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    if f != f:
        return None
    return round(f, nd) if nd is not None else f


def _date(iso):
    try:
        return _dt.date.fromisoformat(str(iso)[:10])
    except (TypeError, ValueError):
        return None


def _head(g: dict, title: str) -> list[list]:
    """The lines on top of every sheet."""
    src = SPOT_WORDS.get(g.get("spot_src") or "", g.get("spot_src") or "")
    day = g.get("on")
    when = ("latest market pass, session %s" % day) if g.get("source") == "latest" else ("closing snapshot of %s" % day)
    return [[title], ["Ticker", g.get("symbol"), g.get("name") or ""], ["Data", when],
            ["Stock price", _num(g.get("spot"), 4), src], ["Options shown", SIDE_WORDS.get(g.get("side") or "", g.get("side"))],
            ["Method", g.get("method") or ""], []]


def listed_contracts(g: dict) -> list[dict]:
    """The listed surface's cells as rows (every cell with an IV)."""
    out, spot, side, ex = [], g.get("spot"), g.get("side") or "otm", g.get("extra") or []
    for r, k in enumerate(g.get("strikes") or ()):
        for c, e in enumerate(g.get("expiries") or ()):
            v = g["cells"][r][c]
            if v is None:
                continue
            x = ex[r][c] if (r < len(ex) and ex[r] and c < len(ex[r])) else None
            out.append({"expiry": _date(e["date"]), "days": e.get("dte"), "monthly": bool(e.get("monthly")),
                        "strike": _num(k), "right": _right(side, float(k), spot), "iv": _num(v, 2),
                        "price": _num((x or {}).get("price"), 4), "last": _num((x or {}).get("last"), 4)})
    return out


def cm_contracts(g: dict) -> list[dict]:
    """The interpolated surface's by-strike cells as rows (strike x horizon)."""
    sg = g.get("by_strike") or g
    on, spot, side = _date(g.get("on")), g.get("spot"), g.get("side") or "otm"
    hz = g.get("horizons") or []
    out = []
    for r, k in enumerate(sg.get("strikes") or ()):
        for c, h in enumerate(hz):
            v = sg["cells"][r][c]
            if v is None:
                continue
            x = sg["extra"][r][c] if r < len(sg.get("extra") or []) else None
            b = sg["brackets"][r][c] if r < len(sg.get("brackets") or []) else None
            near = sg["nearest"][r][c] if r < len(sg.get("nearest") or []) else None
            if b:
                how = "interpolated between %s (%sd) and %s (%sd)" % (b["lo"]["date"], b["lo"]["dte"], b["hi"]["date"],
                                                                       b["hi"]["dte"])
            else:
                how = "the listed expiry %s" % ((near or {}).get("date") or "")
            out.append({"days": h["days"], "date": on + _dt.timedelta(days=int(h["days"])) if on else None,
                        "strike": _num(k), "right": _right(side, float(k), spot), "iv": _num(v, 2),
                        "price": _num((x or {}).get("price"), 4), "how": how})
    return out


def _sheet_rows(ws, rows):
    for row in rows:
        ws.append(row)


def _style(ws, head_n: int, header_row: int, widths: dict):
    from openpyxl.styles import Font  # noqa: PLC0415

    ws.cell(row=1, column=1).font = Font(bold=True, size=13)
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=2)
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def _grid_sheet(wb, name: str, g: dict, title: str, row_label: str, row_keys: list, col_heads: list, cells: list):
    from openpyxl.styles import Font  # noqa: PLC0415
    from openpyxl.utils import get_column_letter  # noqa: PLC0415

    ws = wb.create_sheet(name)
    head = _head(g, title)
    _sheet_rows(ws, head)
    ws.append([row_label] + col_heads)
    hr = len(head) + 1
    for i, rk in enumerate(row_keys):
        ws.append([rk] + [_num(v, 2) for v in cells[i]])
    _style(ws, len(head), hr, {"A": 18})
    for c in range(2, len(col_heads) + 2):
        ws.column_dimensions[get_column_letter(c)].width = 12
    for row in ws.iter_rows(min_row=hr + 1, max_col=1):
        row[0].font = Font(bold=True)
    return ws


def workbook(g: dict, tab: str) -> bytes:
    """The workbook of a surface answer ``g`` (``tab`` ``surface`` or ``cm``) as .xlsx bytes."""
    from openpyxl import Workbook  # noqa: PLC0415

    wb = Workbook()
    ws = wb.active
    ws.title = "Contracts"
    if tab == "cm":
        title = "%s - Interpolated Volatility Surface (fixed horizons)" % g.get("symbol")
        head = _head(g, title)
        _sheet_rows(ws, head)
        ws.append(["Horizon (days)", "Horizon date", "Strike", "Call / put", "IV %", "Price (at that IV)", "How"])
        for r in cm_contracts(g):
            ws.append([r["days"], r["date"], r["strike"], r["right"], r["iv"], r["price"], r["how"]])
        _style(ws, len(head), len(head) + 1, {"A": 15, "B": 13, "C": 10, "D": 18, "E": 8, "F": 17, "G": 62})
        for row in ws.iter_rows(min_row=len(head) + 2, min_col=2, max_col=2):
            row[0].number_format = "yyyy-mm-dd"
        on = _date(g.get("on"))
        heads = ["%dd %s" % (h["days"], (on + _dt.timedelta(days=int(h["days"]))).isoformat() if on else "")
                 for h in g.get("horizons") or []]
        sg = g.get("by_strike") or g
        if g.get("rows"):
            _grid_sheet(wb, "IV grid (moneyness)", g, title, "Strike / price % (strike)",
                        ["%s%% (%s)" % (row["pct"], row["strike"]) for row in g["rows"]], heads, g["cells"])
        _grid_sheet(wb, "IV grid (strike)", g, title, "Strike", [_num(k) for k in sg.get("strikes") or []], heads,
                    sg.get("cells") or [])
    else:
        title = "%s - Volatility Surface (listed expiries)" % g.get("symbol")
        head = _head(g, title)
        _sheet_rows(ws, head)
        ws.append(["Expiry", "Days to expiry", "Monthly", "Strike", "Call / put", "IV %", "Price (at that IV)",
                   "Last trade"])
        for r in listed_contracts(g):
            ws.append([r["expiry"], r["days"], "M" if r["monthly"] else "", r["strike"], r["right"], r["iv"],
                       r["price"], r["last"]])
        _style(ws, len(head), len(head) + 1, {"A": 12, "B": 14, "C": 9, "D": 10, "E": 18, "F": 8, "G": 17, "H": 11})
        for row in ws.iter_rows(min_row=len(head) + 2, max_col=1):
            row[0].number_format = "yyyy-mm-dd"
        heads = ["%s%s %dd" % (e["date"], " M" if e.get("monthly") else "", int(e.get("dte") or 0))
                 for e in g.get("expiries") or []]
        _grid_sheet(wb, "IV grid", g, title, "Strike", [_num(k) for k in g.get("strikes") or []], heads, g.get("cells") or [])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def filename(g: dict, tab: str) -> str:
    sym = re.sub(r"[^A-Za-z0-9]+", "_", str(g.get("symbol") or "surface"))
    kind = "interpolated" if tab == "cm" else "volsurface"
    return "tradehunter-%s-%s-%s.xlsx" % (sym, kind, str(g.get("on") or "")[:10])


__all__ = ["workbook", "filename", "listed_contracts", "cm_contracts"]
